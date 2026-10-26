# -*- coding: utf-8 -*-
"""md → docx 转换(每周上交微云的产物文档用)。

支持: #/##/### 标题、| 表格 |、- 列表、**粗体**、`代码`、> 引用、普通段落。

用法: python scripts/md2docx.py 文档/汇报.md 文档/Week6_李嘉怡_组会汇报.docx
"""
import re
import sys

from docx import Document
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8")

INLINE = [
    (re.compile(r"\*\*(.+?)\*\*"), lambda m: ("bold", m.group(1))),
    (re.compile(r"`(.+?)`"), lambda m: ("code", m.group(1))),
]


def add_inline(p, text):
    pos = 0
    for pat, fn in INLINE:
        for m in pat.finditer(text):
            if m.start() < pos:
                continue
            if m.start() > pos:
                p.add_run(text[pos:m.start()])
            kind, s = fn(m)
            r = p.add_run(s)
            if kind == "bold":
                r.bold = True
            else:
                r.font.name = "Consolas"
                r.font.size = Pt(10)
            pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def md2docx(md_path: str, docx_path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)

    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= {"|", "-", ":", " "}:
            # 表格(单元格列数对齐表头:多则截断,少则补空)
            rows = [[c.strip() for c in r.strip("|").split("|")] for r in (line, *[l for l in lines[i + 2:] if l.startswith("|")])]
            ncols = len(rows[0])
            for k, row in enumerate(rows):
                rows[k] = (row + [""] * ncols)[:ncols]
            tbl = doc.add_table(rows=len(rows), cols=ncols)
            tbl.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    p = tbl.rows[ri].cells[ci].paragraphs[0]
                    add_inline(p, cell)
            i += 2 + (len(rows) - 1)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif line.startswith("> "):
            p = doc.add_paragraph()
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    doc.save(docx_path)
    print(f"已生成: {docx_path}")


if __name__ == "__main__":
    md2docx(sys.argv[1], sys.argv[2])
